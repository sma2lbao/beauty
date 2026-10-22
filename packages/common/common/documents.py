"""附件文本提取(PDF/docx/xlsx/文本):上传时解析,注入 Agent 上下文。

解析失败抛 DocumentParseError,由调用方决定"记录失败但保留附件"还是拒绝。
输出统一截断到 max_chars,避免把整本 PDF 塞进提示词。
"""

from __future__ import annotations

import contextlib
from io import BytesIO
from pathlib import Path

MAX_CHARS = 20_000

_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".json", ".log"}
_PDF_SUFFIXES = {".pdf"}
_DOCX_SUFFIXES = {".docx"}
_XLSX_SUFFIXES = {".xlsx", ".xlsm"}


class DocumentParseError(ValueError):
    """附件无法解析为文本(格式不支持/损坏/加密)。"""


def _suffix(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def extract_text(
    filename: str,
    content_type: str | None,
    data: bytes,
    *,
    max_chars: int = MAX_CHARS,
) -> str:
    suffix = _suffix(filename)
    kind = (content_type or "").split(";", 1)[0].strip().lower()
    if suffix in _TEXT_SUFFIXES or kind.startswith("text/"):
        text = data.decode("utf-8", "replace")
    elif suffix in _PDF_SUFFIXES or kind == "application/pdf":
        text = _pdf_text(data)
    elif suffix in _DOCX_SUFFIXES or kind in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        text = _docx_text(data)
    elif suffix in _XLSX_SUFFIXES or kind in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ):
        text = _xlsx_text(data)
    else:
        raise DocumentParseError(f"不支持解析的附件类型:{suffix or kind or '未知'}")
    text = text.replace("\x00", "").strip()
    if len(text) > max_chars:
        text = text[:max_chars] + "\n…(内容已截断)"
    return text


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 依赖缺失时明确报错
        raise DocumentParseError("PDF 解析依赖未安装") from exc
    try:
        reader = PdfReader(BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"PDF 解析失败:{exc}") from exc
    return "\n\n".join(pages)


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("docx 解析依赖未安装") from exc
    try:
        document = Document(BytesIO(data))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"docx 解析失败:{exc}") from exc
    return "\n".join(lines)


def _xlsx_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("xlsx 解析依赖未安装") from exc
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append("\t".join(values))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"xlsx 解析失败:{exc}") from exc
    finally:
        with contextlib.suppress(Exception):
            workbook.close()  # type: ignore[possibly-undefined]
    return "\n".join(lines)
