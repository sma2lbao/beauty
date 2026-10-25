"""附件文本解析(PDF/docx/xlsx/文本):原 packages/common/tests/test_documents.py 迁入。"""

from __future__ import annotations

from io import BytesIO

import pytest

from agent_orchestrator.services.documents import DocumentParseError, extract_text


def test_text_and_truncation() -> None:
    text = extract_text("note.md", "text/markdown", "你好\n世界".encode())
    assert "你好" in text
    truncated = extract_text("a.txt", "text/plain", ("x" * 100).encode(), max_chars=20)
    assert truncated.endswith("…(内容已截断)")


def test_docx_extraction() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    buffer = BytesIO()
    document.save(buffer)
    text = extract_text(
        "a.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    )
    assert "第一段" in text and "第二段" in text


def test_xlsx_extraction() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报表"
    sheet.append(["姓名", "分数"])
    sheet.append(["小明", 98])
    buffer = BytesIO()
    workbook.save(buffer)
    text = extract_text(
        "a.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        buffer.getvalue(),
    )
    assert "Sheet: 报表" in text
    assert "小明" in text and "98" in text


def test_pdf_extraction_empty_page() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    text = extract_text("a.pdf", "application/pdf", buffer.getvalue())
    assert isinstance(text, str)


def test_unsupported_type() -> None:
    with pytest.raises(DocumentParseError):
        extract_text("a.exe", "application/octet-stream", b"MZ")
